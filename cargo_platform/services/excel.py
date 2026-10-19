"""Чтение и разбор Excel склада (ТЗ, 5.3 и раздел 9).

read_table  — .xlsx / .xls / .csv → список строк-текстов
analyze     — чистая функция: строки + колонки + контекст базы → отчёт до сохранения
save        — сохраняет разобранную партию
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from cargo_platform.db.models import Batch, Parcel, ParcelStatus, Track
from cargo_platform.services.billing import parse_weight
from cargo_platform.services.codes import normalize_code, suggest_codes
from cargo_platform.services.parcels import norm_track

MAX_ROWS = 20000
FIELDS = [
    ("code", "Код клиента"), ("track", "Трек-номер"), ("weight", "Вес (кг)"),
    ("places", "Количество мест"), ("desc", "Описание"), ("volume", "Объёмный вес (кг)"),
]
FIELD_TITLE = dict(FIELDS)
_KEYWORDS = {   # порядок важен: «объёмный вес» проверяем раньше «вес»
    "volume": ("объём", "объем", "volume", "体积"),
    "code": ("код", "code", "клиент", "client", "唛头"),
    "track": ("трек", "track", "накладн", "运单", "tracking"),
    "weight": ("вес", "weight", "кг", "kg", "重量"),
    "places": ("мест", "places", "qty", "количество", "件数"),
    "desc": ("опис", "desc", "товар", "наимен", "name", "品名"),
}


class ExcelError(Exception):
    pass


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_table(data: bytes, filename: str) -> list[list[str]]:
    name = (filename or "").lower()
    rows: list[list[str]] = []
    try:
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            for r in wb.worksheets[0].iter_rows(values_only=True):
                rows.append([_cell(c) for c in r])
                if len(rows) > MAX_ROWS:
                    break
        elif name.endswith(".xls"):
            import xlrd
            sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
            rows = [[_cell(c) for c in sheet.row_values(i)] for i in range(min(sheet.nrows, MAX_ROWS))]
        elif name.endswith(".csv"):
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1251", "replace")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = [[_cell(c) for c in r] for r in csv.reader(io.StringIO(text), dialect)][:MAX_ROWS]
        else:
            raise ExcelError("Пришлите файл .xlsx, .xls или .csv")
    except ExcelError:
        raise
    except Exception as e:
        raise ExcelError("Не получилось прочитать файл. Пришлите .xlsx или .csv") from e
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise ExcelError("Файл пустой.")
    return rows


def guess_mapping(rows: list[list[str]]) -> dict[str, int]:
    """Угадывает колонки по заголовку первой строки: {поле: индекс колонки}."""
    out: dict[str, int] = {}
    for idx, cell in enumerate(rows[0]):
        low = cell.lower()
        for key, words in _KEYWORDS.items():
            if key not in out and any(w in low for w in words):
                out[key] = idx
                break
    return out


def detect_header(rows: list[list[str]], mapping: dict[str, int]) -> bool:
    """Первая строка — заголовок: в ней находятся названия колонок или вместо веса стоит текст."""
    first = rows[0]
    if len(guess_mapping(rows)) >= 1 and not any(c.replace(",", ".").replace(".", "").isdigit() for c in first if c):
        return True
    w = mapping.get("weight")
    if w is not None and w < len(first) and first[w] and parse_weight(first[w]) is None:
        return True
    return False


def fit_mapping(mapping: dict[str, int], ncols: int) -> dict[str, int]:
    """Файл склада стал уже обычного: необязательные колонки за его краем отбрасываем."""
    return {k: v for k, v in mapping.items() if 0 <= v < ncols}


def valid_mapping(mapping: dict[str, int], ncols: int) -> bool:
    return (("code" in mapping or "track" in mapping)
            and all(0 <= i < ncols for i in mapping.values()))


@dataclass
class Row:
    line: int
    raw_code: str
    code: str | None
    track: str | None
    weight: Decimal | None
    places: int
    desc: str
    volume: Decimal | None
    client_id: int | None = None
    kind: str = "ok"                 # ok | unmatched | dup
    note: str = ""
    suggestions: list[str] = field(default_factory=list)
    existing_parcel_id: int | None = None
    merged: int = 0                  # сколько строк того же трека слито в эту


@dataclass
class Ctx:
    prefix: str
    digits: int
    next_number: int
    clients_by_code: dict[str, int]
    client_by_track: dict[str, int]
    parcels_by_track: dict[str, tuple[int, int | None, ParcelStatus]]   # id, client_id, status


@dataclass
class Analysis:
    rows: list[Row]
    skipped_empty: int = 0

    @property
    def stored(self) -> list[Row]:
        return [r for r in self.rows if r.kind != "dup"]

    @property
    def found(self) -> list[Row]:
        return [r for r in self.rows if r.kind == "ok"]

    @property
    def unmatched(self) -> list[Row]:
        return [r for r in self.rows if r.kind == "unmatched"]

    @property
    def dups(self) -> list[Row]:
        return [r for r in self.rows if r.kind == "dup"]

    @property
    def client_ids(self) -> set[int]:
        return {r.client_id for r in self.found if r.client_id}

    @property
    def weight_total(self) -> Decimal:
        return sum((r.weight or Decimal(0) for r in self.found), Decimal(0))

    @property
    def merged(self) -> int:
        return sum(r.merged for r in self.rows)

    @property
    def no_weight(self) -> int:
        return sum(1 for r in self.stored if r.weight is None)


def _int(raw, default=1) -> int:
    try:
        return max(1, min(999, int(float(str(raw).replace(",", ".")))))
    except (TypeError, ValueError):
        return default


def analyze(rows: list[list[str]], mapping: dict[str, int], header: bool, ctx: Ctx) -> Analysis:
    def get(r: list[str], key: str) -> str:
        i = mapping.get(key)
        return r[i].strip() if i is not None and i < len(r) else ""

    data = rows[1:] if header else rows
    result: list[Row] = []
    by_track: dict[str, Row] = {}
    codes = list(ctx.clients_by_code)
    for n, r in enumerate(data, start=2 if header else 1):
        raw_code, raw_track = get(r, "code"), get(r, "track")
        track = norm_track(raw_track, strict=False)
        if not raw_code and not track:
            continue
        weight = parse_weight(get(r, "weight") or None)
        places = _int(get(r, "places")) if get(r, "places") else 1
        row = Row(line=n, raw_code=raw_code, code=None, track=track, weight=weight, places=places,
                  desc=get(r, "desc")[:200], volume=parse_weight(get(r, "volume") or None))
        # одна посылка пришла двумя местами внутри файла → одна запись, места и вес суммируются
        if track and track in by_track:
            first = by_track[track]
            first.places += row.places
            if row.weight is not None:
                first.weight = (first.weight or Decimal(0)) + row.weight
            if row.volume is not None:
                first.volume = (first.volume or Decimal(0)) + row.volume
            first.merged += 1
            continue
        if track:
            by_track[track] = row
        result.append(row)

    for row in result:
        # дубль: этот трек уже принимался раньше — повторно не начисляем
        known = ctx.parcels_by_track.get(row.track) if row.track else None
        if known and known[2] != ParcelStatus.expected:
            row.kind, row.note, row.client_id = "dup", "уже есть в базе", known[1]
            continue
        if known:
            row.existing_parcel_id = known[0]
        # клиент по коду
        code, foreign_prefix = normalize_code(row.raw_code, ctx.prefix, ctx.digits) if row.raw_code else (None, False)
        row.code = code
        if code and code in ctx.clients_by_code:
            row.client_id = ctx.clients_by_code[code]
            continue
        # кода нет или он не нашёлся — ищем по треку среди добавленных клиентами
        by_t = ctx.client_by_track.get(row.track) if row.track else None
        if by_t is None and known and known[1]:
            by_t = known[1]
        if by_t is not None:
            row.client_id = by_t
            continue
        row.kind = "unmatched"
        if not row.raw_code:
            row.note = "нет кода, трек не добавлен клиентом"
        elif foreign_prefix:
            row.note = "код с ошибкой (чужой префикс)"
        elif code is None:
            row.note = "код не распознан"
        else:
            number = int(code.split("-")[1])
            row.note = "клиент удалён" if number < ctx.next_number else "такой код не выдан"
        row.suggestions = suggest_codes(row.raw_code, codes)
    return Analysis(rows=result)


async def load_ctx(s: AsyncSession, company_id: int, prefix: str, digits: int, next_number: int,
                   clients: dict[str, int], tracks: list[str]) -> Ctx:
    """Читает из базы только то, что нужно файлу: клиентов и треки из файла."""
    by_track: dict[str, int] = {}
    parcels: dict[str, tuple[int, int | None, ParcelStatus]] = {}
    uniq = list({t for t in tracks if t})
    for i in range(0, len(uniq), 500):
        chunk = uniq[i:i + 500]
        for t, cid in (await s.execute(select(Track.track_number, Track.client_id).where(
                Track.company_id == company_id, Track.track_number.in_(chunk)))).all():
            by_track[t] = cid
        for pid, t, cid, st in (await s.execute(select(Parcel.id, Parcel.track_number, Parcel.client_id, Parcel.status).where(
                Parcel.company_id == company_id, Parcel.track_number.in_(chunk)))).all():
            parcels[t] = (pid, cid, st)
    return Ctx(prefix=prefix, digits=digits, next_number=next_number, clients_by_code=clients,
               client_by_track=by_track, parcels_by_track=parcels)


def file_tracks(rows: list[list[str]], mapping: dict[str, int], header: bool) -> list[str]:
    i = mapping.get("track")
    if i is None:
        return []
    out = []
    for r in (rows[1:] if header else rows):
        if i < len(r) and (t := norm_track(r[i], strict=False)):
            out.append(t)
    return out


async def save(s: AsyncSession, company_id: int, batch: Batch, analysis: Analysis,
               status: ParcelStatus = ParcelStatus.china_warehouse) -> list[Parcel]:
    """Сохраняет посылки партии. Неопознанные тоже сохраняются (client_id пусто)."""
    saved: list[Parcel] = []
    for row in analysis.stored:
        fields = dict(batch_id=batch.id, client_id=row.client_id, raw_code=row.raw_code or None,
                      track_number=row.track, weight=row.weight, volume_weight=row.volume,
                      places=row.places, description=row.desc or None,
                      problem_note=row.note if row.kind == "unmatched" else None, status=status)
        if row.existing_parcel_id:
            p = await s.get(Parcel, row.existing_parcel_id)
            for k, v in fields.items():
                if k == "description" and v is None:
                    continue
                setattr(p, k, v)
        else:
            p = Parcel(company_id=company_id, **fields)
            s.add(p)
        saved.append(p)
    batch.places = sum(r.places for r in analysis.stored)
    batch.total_weight = sum((r.weight or Decimal(0) for r in analysis.stored), Decimal(0))
    await s.flush()
    return saved
